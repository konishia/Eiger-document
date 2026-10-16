"""Excel ファイル（.xlsx / .xlsm / .xls）を表（行 × 列の文字列）として読み取る

従業員名簿の一括読み込みで使う。列の意味の判定や内容の確認は画面側で行うので、ここでは
「データのある最初のシート」を見たままの文字列にして返すだけにする。
"""

import base64
import binascii
import io
import re
from datetime import date, datetime, time

import openpyxl
import xlrd
from fastapi import HTTPException

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLS = 50

# ファイル先頭のバイト列で形式を判定する（拡張子は信用しない）
_XLSX_SIGNATURE = b"PK\x03\x04"  # ZIP 形式（.xlsx / .xlsm）
_XLS_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # 旧形式（.xls）


def _cell_text(value: object, number_format: str = "General") -> str:
    """セルの値を、Excel で見えている文字に近い文字列にする"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d") if value.time() == time(0) else value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        # 「00123」のように 0 埋めの表示形式なら、先頭の 0 を戻す（従業員コードでよくある）
        zeros = re.fullmatch(r"0+", number_format or "")
        return str(value).zfill(len(zeros.group())) if zeros else str(value)
    return str(value).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """末尾の空の列・空の行を落とす"""
    rows = [row[:MAX_COLS] for row in rows]
    rows = [row[: max((i + 1 for i, c in enumerate(row) if c), default=0)] for row in rows]
    while rows and not rows[-1]:
        rows.pop()
    return rows


def _read_xlsx(data: bytes) -> list[tuple[str, list[list[str]]]]:
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = []
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(max_row=MAX_ROWS + 1, max_col=MAX_COLS):
                rows.append([_cell_text(c.value, getattr(c, "number_format", "General")) for c in row])
            sheets.append((sheet.title, _trim(rows)))
        return sheets
    finally:
        book.close()


def _read_xls(data: bytes) -> list[tuple[str, list[list[str]]]]:
    book = xlrd.open_workbook(file_contents=data, formatting_info=False)
    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(min(sheet.nrows, MAX_ROWS + 1)):
            cells = []
            for c in range(min(sheet.ncols, MAX_COLS)):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    cells.append(_cell_text(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    cells.append(_cell_text(bool(cell.value)))
                else:
                    cells.append(_cell_text(cell.value))
            rows.append(cells)
        sheets.append((sheet.name, _trim(rows)))
    return sheets


def read_first_sheet(encoded: str) -> tuple[str, list[list[str]]]:
    """Base64 の Excel ファイルから、データのある最初のシートを (シート名, 表) で返す。読めなければ 400 エラー"""
    try:
        data = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as e:
        raise HTTPException(status_code=400, detail="ファイルのデータが壊れています") from e
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="ファイルが大きすぎます（2MBまで）")

    try:
        if data.startswith(_XLSX_SIGNATURE):
            sheets = _read_xlsx(data)
        elif data.startswith(_XLS_SIGNATURE):
            sheets = _read_xls(data)
        else:
            raise HTTPException(status_code=400, detail="Excelのファイル（.xlsx / .xls）を選んでください")
    except HTTPException:
        raise
    except Exception as e:
        # パスワード付き・壊れたファイルなど
        raise HTTPException(
            status_code=400,
            detail="Excelのファイルを読み取れませんでした（パスワード付きのファイルは、パスワードを外してから選んでください）",
        ) from e

    for name, rows in sheets:
        if any(any(cell for cell in row) for row in rows):
            if len(rows) > MAX_ROWS:
                raise HTTPException(status_code=400, detail=f"行が多すぎます（{MAX_ROWS}行まで）")
            return name, rows
    raise HTTPException(status_code=400, detail="Excelのファイルにデータがありません")
